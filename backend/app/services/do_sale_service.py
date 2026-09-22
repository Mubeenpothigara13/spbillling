from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Optional

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib.units import mm
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.audit import AuditAction
from app.models.bill import Bill
from app.models.customer import Customer
from app.models.do_sale import DoSale
from app.models.product import ProductVariant
from app.models.user import User
from app.schemas.do_sale import DoSaleCreate, DoSaleSummary, DoSaleVariantTotal
from app.services.report_service import (
    _HEADER_FILL, _HEADER_FONT, _autofit, _fmt_date_str, _header_block,
    _make_table, _new_doc, _subtitle_para,
)
from app.utils.audit import write_audit


def create_sales(db: Session, user: User, payload: DoSaleCreate) -> list[DoSale]:
    if user.do_id is None:
        raise HTTPException(
            status_code=403, detail="Only a distributor outlet login can record a sale"
        )
    sale_date = payload.sale_date or date.today()

    rows: list[DoSale] = []
    for line in payload.lines:
        customer = db.get(Customer, line.customer_id)
        # A DO can only sell to its own customers.
        if not customer or customer.is_deleted or customer.do_id != user.do_id:
            raise HTTPException(status_code=400, detail="Customer not found")
        variant = db.get(ProductVariant, line.product_variant_id)
        if not variant or not variant.is_active:
            raise HTTPException(status_code=400, detail="Product not found")
        rows.append(DoSale(
            do_id=user.do_id,
            customer_id=customer.id,
            sale_date=sale_date,
            product_variant_id=variant.id,
            quantity=line.quantity,
            rate=line.rate if line.rate is not None else variant.unit_price,
            gst_rate=variant.gst_rate,
            empty_returned=line.empty_returned,
            created_by_id=user.id,
        ))
    db.add_all(rows)
    db.flush()
    write_audit(db, entity_type="do_sale", entity_id=rows[0].id,
                action=AuditAction.CREATE, user_id=user.id,
                changes={"lines": len(rows), "sale_date": sale_date.isoformat(),
                         "quantity": sum(r.quantity for r in rows)})
    ids = [r.id for r in rows]
    db.commit()
    return list(db.scalars(
        select(DoSale).where(DoSale.id.in_(ids)).order_by(DoSale.id)
    ).unique())


def _conditions(
    *,
    do_id: Optional[int],
    status: Optional[str] = None,
    from_date: Optional[date] = None,
    to_date: Optional[date] = None,
    customer_id: Optional[int] = None,
) -> list:
    conds = []
    if do_id is not None:
        conds.append(DoSale.do_id == do_id)
    if status == "pending":
        conds.append(DoSale.bill_id.is_(None))
    elif status == "billed":
        conds.append(DoSale.bill_id.is_not(None))
    if from_date is not None:
        conds.append(DoSale.sale_date >= from_date)
    if to_date is not None:
        conds.append(DoSale.sale_date <= to_date)
    if customer_id is not None:
        conds.append(DoSale.customer_id == customer_id)
    return conds


def list_sales(**filters):
    return (
        select(DoSale)
        .where(*_conditions(**filters))
        .order_by(DoSale.sale_date.desc(), DoSale.id.desc())
    )


def summary(
    db: Session, *, do_id: Optional[int], from_date: Optional[date], to_date: Optional[date]
) -> DoSaleSummary:
    conds = _conditions(do_id=do_id, from_date=from_date, to_date=to_date)
    qty = func.sum(DoSale.quantity)
    by_variant = db.execute(
        select(ProductVariant.name, qty)
        .select_from(DoSale)
        .join(ProductVariant, ProductVariant.id == DoSale.product_variant_id)
        .where(*conds)
        .group_by(ProductVariant.name)
        .order_by(qty.desc())
    ).all()
    customers = db.scalar(
        select(func.count(func.distinct(DoSale.customer_id))).where(*conds)
    ) or 0
    totals = [DoSaleVariantTotal(variant_name=n, qty=int(q or 0)) for n, q in by_variant]
    return DoSaleSummary(
        total_qty=sum(t.qty for t in totals),
        customer_count=customers,
        by_variant=totals,
    )


def quantities_by_variant_name(db: Session, do_id: Optional[int]) -> list[tuple[str, int]]:
    """All-time quantity sold per variant name — feeds the Indent stock."""
    qty = func.sum(DoSale.quantity)
    rows = db.execute(
        select(ProductVariant.name, qty)
        .select_from(DoSale)
        .join(ProductVariant, ProductVariant.id == DoSale.product_variant_id)
        .where(*_conditions(do_id=do_id))
        .group_by(ProductVariant.name)
    ).all()
    return [(name, int(q or 0)) for name, q in rows]


def link_to_bill(db: Session, sale_ids: list[int], bill: Bill) -> None:
    """Mark pending DO sale lines as billed by `bill`. Every line must still
    be pending and belong to the bill's customer."""
    wanted = set(sale_ids)
    rows = list(db.scalars(select(DoSale).where(DoSale.id.in_(wanted))).unique())
    if len(rows) != len(wanted):
        raise HTTPException(status_code=400, detail="DO sale not found")
    for r in rows:
        if r.bill_id is not None:
            raise HTTPException(status_code=400, detail="A DO sale is already billed")
        if r.customer_id != bill.customer_id:
            raise HTTPException(
                status_code=400, detail="A DO sale belongs to a different customer"
            )
        r.bill_id = bill.id


# ---------- DO's Report export (a DO's own sales — no bill numbers) ----------

def _report_rows(db: Session, *, do_id: Optional[int], from_date: date, to_date: date) -> list[DoSale]:
    return list(db.scalars(
        select(DoSale)
        .where(*_conditions(do_id=do_id, from_date=from_date, to_date=to_date))
        .order_by(DoSale.sale_date.asc(), DoSale.id.asc())
    ).unique())


def report_excel(db: Session, *, do_id: Optional[int], from_date: date, to_date: date) -> bytes:
    rows = _report_rows(db, do_id=do_id, from_date=from_date, to_date=to_date)
    wb = Workbook()
    ws = wb.active
    ws.title = "DO's Report"
    ws.append([f"DO's Report · {from_date.strftime('%d %b %Y')} → {to_date.strftime('%d %b %Y')}"])
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=6)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    headers = ["Date", "Customer", "Mobile", "Product", "Qty", "Status"]
    ws.append([])
    ws.append(headers)
    for cell in ws[3]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    total_qty = 0
    for r in rows:
        ws.append([
            r.sale_date.strftime("%Y-%m-%d"), r.customer_name, r.customer_mobile or "-",
            r.variant_name, r.quantity, "Billed" if r.bill_id else "Pending",
        ])
        total_qty += r.quantity
    ws.append([])
    ws.append(["TOTAL", "", "", "", total_qty, ""])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    _autofit(ws)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_pdf(db: Session, *, do_id: Optional[int], from_date: date, to_date: date) -> bytes:
    rows = _report_rows(db, do_id=do_id, from_date=from_date, to_date=to_date)
    subtitle = (
        f"{from_date.strftime('%d %b %Y')} → {to_date.strftime('%d %b %Y')}"
        f" · {len(rows)} sale{'s' if len(rows) != 1 else ''}"
    )
    buf = BytesIO()
    doc = _new_doc(buf)
    story: list = _header_block("DO's Report", subtitle)
    if not rows:
        story.append(_subtitle_para("No sales in the selected range."))
        doc.build(story)
        return buf.getvalue()
    data = [
        [_fmt_date_str(r.sale_date.isoformat()), r.customer_name, r.customer_mobile or "-",
         r.variant_name, str(r.quantity), "Billed" if r.bill_id else "Pending"]
        for r in rows
    ]
    total_qty = sum(r.quantity for r in rows)
    totals = ["TOTAL", "", "", "", str(total_qty), ""]
    col_widths = [26 * mm, 34 * mm, 26 * mm, 34 * mm, 16 * mm, 20 * mm]
    story.append(_make_table(
        ["Date", "Customer", "Mobile", "Product", "Qty", "Status"],
        data,
        col_widths=col_widths,
        totals=totals,
        right_align_cols=[4],
        center_align_cols=[5],
    ))
    doc.build(story)
    return buf.getvalue()
